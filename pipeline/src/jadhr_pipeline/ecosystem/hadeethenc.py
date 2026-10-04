"""HadeethEnc's explanations, benefits and word meanings (PLAN P11.2c, D8, research 21 §3): the Arabic export v1.7.0
verbatim — the publisher's terms forbid any modification, addition or deletion and ask for the version — and the
reference each item gives (from the API, the export lacking it), read for the Six Books' numbers it cites.

An item is linked to the hadith it cites by number in our numbering (al-Bukhārī Fatḥ al-Bārī's, Muslim ʿAbd al-Bāqī's,
the Sunan their printings'), with how much of its wording our record holds. It is HadeethEnc's item citing that number,
never our text: its hadith is shown as HadeethEnc gives it, its grade as HadeethEnc's verdict. A citation whose wording
our record does not hold is kept for review, not shown with the hadith."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from jadhr_arabic import normalize

VERSION = "1.7.0"
# the Six Books as the references name them, and our texts
BOOKS = {
    "صحيح البخاري": "bukhari",
    "صحيح مسلم": "muslim",
    "سنن أبي داود": "abudawud",
    "سنن الترمذي": "tirmidhi",
    "جامع الترمذي": "tirmidhi",
    "سنن النسائي": "nasai",
    "السنن الصغرى للنسائي": "nasai",
    "سنن ابن ماجه": "ibnmajah",
}
_W = re.compile(r"[ء-غف-ي]+")
AGREES = 0.6
PARTLY = 0.3


def items(raw: Path, spec: dict[str, Any]) -> list[dict[str, Any]]:
    """The export's rows, verbatim, each with the reference the API gives it."""
    import openpyxl

    wb = openpyxl.load_workbook(raw / spec["excel"], read_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    wb.close()
    head = [str(c) for c in rows[1]]
    if head != [
        "id",
        "title",
        "hadith_text",
        "explanation",
        "word_meanings",
        "benefits",
        "grade",
        "takhrij",
        "link",
    ]:
        raise ValueError(f"HadeethEnc export columns changed: {head}")
    api = {x["id"]: x for x in json.loads((raw / spec["api"]).read_text(encoding="utf-8"))}
    out = []
    for r in rows[2:]:
        if r[0] is None:
            continue
        rec = dict(zip(head, r, strict=True))
        a = api.get(str(rec["id"])) or {}
        out.append(
            {
                "id": int(rec["id"]),
                "title": rec["title"] or "",
                "text": rec["hadith_text"] or "",
                "explanation": rec["explanation"] or "",
                "word_meanings": _meanings(rec["word_meanings"]),
                "benefits": [b for b in (rec["benefits"] or "").split("\n") if b.strip()],
                "grade": rec["grade"],
                "attribution": rec["takhrij"],
                "link": rec["link"],
                "reference": a.get("reference"),
                "version": VERSION,
            }
        )
    return out


def _meanings(text: str | None) -> list[dict[str, str]]:
    """«حجة الوداع : سُمِّيَتْ …» a line each, the word and its meaning as the export writes them; «....» is none."""
    out = []
    for line in (text or "").split("\n"):
        if not line.strip() or line.strip(". ") == "":
            continue
        word, sep, meaning = line.partition(":")
        out.append(
            {"word": word.strip(), "meaning": meaning.strip()}
            if sep
            else {"word": "", "meaning": line.strip()}
        )
    return out


_BOOK = re.compile("|".join(sorted((re.escape(b) for b in BOOKS), key=len, reverse=True)))


def cited(reference: str | None) -> list[tuple[str, int]]:
    """The Six Books' numbers a reference cites, each to the book named in the words just before it: «صحيح مسلم
    (2/ 646) (939).» → (muslim, 939); «صحيح البخاري (9/ 22) (6953)، (1/ 6) (1).» → 6953 and 1, only punctuation
    between; «سنن أبي داود (…) (4607)، وسنن الترمذي (…) (2676)» → each its own; «… وصحيح ابن حبان (97)» → none, the
    words before it naming a book outside the Six. A volume/page group is passed over; a short range gives both ends
    and between."""
    out: list[tuple[str, int]] = []
    for line in (reference or "").split("\n"):
        book: str | None = None
        at = 0
        for m in re.finditer(r"\(([^)]*)\)", line):
            words = line[at : m.start()]
            at = m.end()
            if _W.search(words):
                named = list(_BOOK.finditer(words))
                book = BOOKS[named[-1].group()] if named else None
            group = m.group(1)
            if book is None or "/" in group or "ص" in group or not re.search(r"\d", group):
                continue
            nums = [int(n) for n in re.findall(r"\d+", group)]
            if len(nums) == 2 and re.search(r"[-–]", group) and 0 < nums[1] - nums[0] <= 3:
                nums = list(range(nums[0], nums[1] + 1))
            out += [(book, n) for n in nums]
    return list(dict.fromkeys(out))


def _words(text: str) -> set[str]:
    return set(_W.findall(normalize(text)))


def agreement(item_text: str, record_text: str) -> float:
    """How much of the item's wording our record holds: the share of its words found in ours."""
    theirs = _words(item_text)
    return len(theirs & _words(record_text)) / len(theirs) if theirs else 0.0


def links(
    all_items: list[dict[str, Any]], records: dict[tuple[str, str], tuple[int, str]]
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    """Each item to the records it cites: ``records`` maps (book, our ref) to (passage id, text). Muslim's
    ʿAbd al-Bāqī number covers its sub-records («939.01», «939.02»), the one whose wording agrees most is taken."""
    by_number: dict[tuple[str, int], list[tuple[int, str]]] = {}
    for (book, ref), rec in records.items():
        try:
            n = int(float(ref)) if book == "muslim" else int(ref)
        except ValueError:
            continue
        by_number.setdefault((book, n), []).append(rec)
    out = []
    stats = {"cited": 0, "agrees": 0, "partly": 0, "disagrees": 0, "absent": 0}
    for it in all_items:
        for book, n in cited(it["reference"]):
            stats["cited"] += 1
            cands = by_number.get((book, n), [])
            if not cands:
                stats["absent"] += 1
                continue
            score, pid = max((agreement(it["text"], text), pid) for pid, text in cands)
            status = "agrees" if score >= AGREES else "partly" if score >= PARTLY else "disagrees"
            stats[status] += 1
            out.append(
                {
                    "item_id": it["id"],
                    "passage_id": pid,
                    "book": book,
                    "number": n,
                    "agreement": round(score, 3),
                    "status": status,
                }
            )
    return out, stats
