"""The narrator linking measured against MIS / MIDS for Ṣaḥīḥ Muslim (Mendeley gzprcr93zn v2, CC BY 4.0): a
cross-check, never shown. MIS maps each chain of each hadith to its narrators by hand, with their Arabic names and,
for most, a ṭabaqa beside the English name («Sufyan bin Sa‘id Ath-Thawri (7)»).

A Muslim record joins an MIS hadith when their texts are equal once normalised. Then: of the names we linked, how many
are one of the narrators MIS gives the hadith (precision); of MIS's narrators, how many we reached (recall); and where
both give a ṭabaqa, whether they agree. Two names are the same narrator when their first words agree — the name and
the father's («عروه بن الزبير»), or the kunya («ابو هريره»).

Needs openpyxl (the workspace's dev group); the build never runs it."""

from __future__ import annotations

import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from jadhr_arabic import normalize

from .narrators import Link, canon

MIS = "mis-muslim-narrators"
_TABAQA = re.compile(r"\((\d+)\)\s*$")
_W = re.compile(r"[ء-غف-ي]+")


def _sheet(path: Path) -> list[tuple[Any, ...]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    wb.close()
    return rows[1:]


def _key(text: str) -> str:
    return " ".join(_W.findall(normalize(text or "")))


def mis(raw: Path) -> tuple[dict[str, dict[str, Any]], dict[str, list[str]]]:
    """MIS's narrators (id → names, ṭabaqa) and each hadith's narrators by its text."""
    nodes = {}
    for label, nid, name_ar, mapped_ar, _en, _g, _gen in _sheet(
        raw / MIS / "2_Hadith_SahihMuslim_Narrators=Nodes for Graph.xlsx"
    ):
        m = _TABAQA.search(str(label or ""))
        nodes[str(nid)] = {
            "names": [n for n in (name_ar, mapped_ar) if n],
            "tabaqa": int(m.group(1)) if m and m.group(1) != "0" else (1 if m else None),
        }
    texts = {str(r[2]): _key(r[8]) for r in _sheet(raw / MIS / "1_Hadith_SahihMuslim_HadithContent.xlsx")}
    by_hadith: dict[str, set[str]] = defaultdict(set)
    for r in _sheet(raw / MIS / "3_Hadith_SahihMuslim_Isnad=Edges for Graph.xlsx"):
        by_hadith[str(r[1])] |= {str(r[5]), str(r[9])}
    by_text: dict[str, list[str]] = {}
    for h, ids in by_hadith.items():
        if texts.get(h):
            by_text[texts[h]] = sorted(ids)
    return nodes, by_text


_TITLES = ("ام المومنين", "امير المومنين")


def _units(words: list[str]) -> list[str]:
    """A name's units: «عبد الرزاق» one, the article off («العباس» is «عباس»), a title before it dropped."""
    s = " ".join(words)
    for t in _TITLES:
        s = s.removeprefix(t + " ")
    out: list[str] = []
    for w in s.split():
        if out and out[-1] == "عبد":
            out[-1] = "عبد " + w
            continue
        out.append(w.removeprefix("ال") if len(w) > 4 and w != "الله" else w)
    return out


def same(a: list[str], b: list[str]) -> bool:
    """The same narrator by name: the kunya, or the name with the father's (the first three units) — at the start of
    both names, or the one's within the other («ابو مرثد الغنوي كناز بن الحصين» holds «كناز بن الحصين»)."""
    a, b = _units(a), _units(b)
    if not a or not b:
        return False
    if a[0] == "ابو" and b[0] == "ابو":
        return a[:2] == b[:2]
    if len(a) >= 3 and len(b) >= 3 and a[1] in ("بن", "بنت") and b[1] in ("بن", "بنت") and a[:3] == b[:3]:
        return True

    def within(x: list[str], y: list[str]) -> bool:
        head = x[:3] if len(x) >= 3 and x[1] in ("بن", "بنت") else x[:2] if x[0] == "ابو" else x[:1]
        n = len(head)
        return n >= 2 and any(y[k : k + n] == head for k in range(len(y) - n + 1))

    return within(a, b) or within(b, a) or (a[0] == b[0] and (len(a) == 1 or len(b) == 1))


def measure(
    raw: Path,
    records: list[dict[str, Any]],
    links: list[list[Link] | list[dict[str, Any]]],
    entries: dict[int, Any],
) -> dict[str, Any]:
    """Our Muslim links against MIS's: precision, recall, ṭabaqa agreement, over the records that join."""
    nodes, by_text = mis(raw)
    c: Counter[str] = Counter()
    for h, ls in zip(records, links, strict=True):
        ids = by_text.get(_key(h["text"]))
        if ids is None:
            continue
        c["records"] += 1
        theirs = [(i, [canon(n) for n in nodes[i]["names"]]) for i in ids if i in nodes]
        reached: set[str] = set()
        for link in ls:
            chosen, rule = (
                (link["chosen"], link["rule"]) if isinstance(link, dict) else (link.chosen, link.rule)
            )
            if not chosen:
                continue
            e = entries[chosen]
            ours = canon(e.name)
            match = next((i for i, names in theirs if any(same(ours, n) for n in names)), None)
            c["linked"] += 1
            c[f"rule {rule}"] += 1
            if match is not None:
                c["linked and in MIS"] += 1
                c[f"rule {rule} in MIS"] += 1
                reached.add(match)
                t = nodes[match]["tabaqa"]
                if t is not None and e.tabaqa is not None:
                    c["tabaqa compared"] += 1
                    c["tabaqa same"] += t == e.tabaqa
        c["MIS narrators"] += len(theirs)
        c["MIS narrators reached"] += len(reached)
    by_rule = {
        k.removeprefix("rule "): round(c[f"{k} in MIS"] / v, 3)
        for k, v in c.items()
        if k.startswith("rule ") and not k.endswith(" in MIS") and v >= 20
    }
    return {
        **{k: v for k, v in c.items() if not k.startswith("rule ")},
        "precision by rule": by_rule,
        "precision": round(c["linked and in MIS"] / max(c["linked"], 1), 3),
        "recall": round(c["MIS narrators reached"] / max(c["MIS narrators"], 1), 3),
        "tabaqa agreement": round(c["tabaqa same"] / max(c["tabaqa compared"], 1), 3),
    }
